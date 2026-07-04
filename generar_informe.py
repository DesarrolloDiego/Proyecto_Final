from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as RLImage,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from perurecetas_core import catalogar_dataset


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
ASSETS_DIR = OUTPUT_DIR / "assets"
DOCX_PATH = OUTPUT_DIR / "informe_perurecetas.docx"
PDF_PATH = OUTPUT_DIR / "informe_perurecetas.pdf"
ARCH_PATH = ASSETS_DIR / "arquitectura_perurecetas.png"


PROJECT_TITLE = "Agente de Recetas y Nutricion Peruana con RAG, FAISS y API externa"
PROJECT_NAME = "PeruRecetas"
INTEGRANTES = [
    "GIOVANNI FRANCO LUCA BARCENA ALDAVE",
    "DIEGO ALONSO MERINO CASTAÑEDA",
]
DOCENTE = "EDUARDO YAURI LOZANO"
CURSO = "Redes Neuronales Artificiales"
FECHA = "04/07/2026"


def ensure_dirs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)


def crear_diagrama() -> None:
    ensure_dirs()
    width, height = 1400, 760
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    try:
        font_title = ImageFont.truetype("arial.ttf", 34)
        font_box = ImageFont.truetype("arial.ttf", 22)
        font_small = ImageFont.truetype("arial.ttf", 18)
    except Exception:
        font_title = ImageFont.load_default()
        font_box = ImageFont.load_default()
        font_small = ImageFont.load_default()

    draw.text((50, 35), "Arquitectura de PeruRecetas", fill=(15, 37, 69), font=font_title)

    boxes = [
        ((70, 140, 330, 260), "Usuario\nNotebook / Streamlit"),
        ((430, 90, 710, 210), "Agente\nMemoria + Orquestacion"),
        ((830, 80, 1220, 200), "Gemini LLM\nRespuesta final"),
        ((430, 300, 710, 430), "Tools Python\nRecetas y sustituciones"),
        ((830, 290, 1220, 420), "API externa\nComida peruana por region"),
        ((430, 520, 710, 660), "TXT etiquetados\nDataset RAG"),
        ((830, 520, 1220, 660), "FAISS + Embeddings\nRecuperacion semantica"),
    ]
    for rect, label in boxes:
        draw.rounded_rectangle(rect, radius=18, outline=(46, 116, 181), width=3, fill=(242, 246, 249))
        lines = label.split("\n")
        y = rect[1] + 28
        for line in lines:
            draw.text((rect[0] + 24, y), line, fill=(20, 30, 40), font=font_box if y == rect[1] + 28 else font_small)
            y += 32

    arrows = [
        ((330, 200), (430, 150)),
        ((710, 150), (830, 140)),
        ((570, 210), (570, 300)),
        ((710, 365), (830, 355)),
        ((570, 520), (570, 430)),
        ((710, 590), (830, 590)),
        ((1025, 520), (1025, 420)),
        ((1025, 290), (1025, 200)),
    ]
    for start, end in arrows:
        draw.line([start, end], fill=(31, 78, 121), width=4)
        x, y = end
        draw.polygon([(x, y), (x - 12, y - 8), (x - 8, y + 12)], fill=(31, 78, 121))

    img.save(ARCH_PATH)


def set_run(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, color="2E74B5" if level <= 2 else "1F4D78")
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def crear_docx() -> None:
    crear_diagrama()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(PROJECT_TITLE)
    set_run(run, size=20, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(PROJECT_NAME)
    set_run(run, size=16, bold=True, color="2E74B5")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Curso: {CURSO}\n").bold = True
    meta.add_run(f"Docente: {DOCENTE}\n")
    meta.add_run("Integrantes:\n")
    for integrante in INTEGRANTES:
        meta.add_run(f"- {integrante}\n")
    meta.add_run(f"Fecha de entrega: {FECHA}")
    doc.add_page_break()

    add_heading(doc, "A. Descripción general de los modelos o librerías utilizadas", 1)
    add_bullets(
        doc,
        [
            "Gemini (google-generativeai): Modelo LLM principal usado para la generación de respuestas y embeddings.",
            "AIsuite: Cliente unificado que permite interactuar con múltiples proveedores de LLMs. Se configuró para interactuar con Gemini.",
            "FAISS (faiss-cpu): Librería de Meta para búsqueda de similitud eficiente y agrupamiento de vectores densos. Usada para el índice del RAG.",
            "Streamlit: Framework para crear aplicaciones web interactivas en Python, usado para la UI del proyecto.",
            "requests / urllib: Utilizados para las llamadas HTTP a la API externa de comida peruana.",
            "pdfplumber / pypdf: Utilizados inicialmente para extraer texto de PDFs (ahora se usan TXTs extraídos).",
        ],
    )

    add_heading(doc, "B. Flujo de trabajo personalizado", 1)
    doc.add_paragraph(
        "El flujo de trabajo comienza cuando el usuario introduce una consulta a través de la interfaz web (Streamlit). "
        "El Agente recibe la consulta e intenta enriquecer el contexto mediante dos vías paralelas: "
        "1. Realiza una búsqueda semántica en la base de datos vectorial (FAISS) para recuperar recetas locales relevantes (RAG). "
        "2. Identifica si la consulta requiere información de la API externa (por ejemplo, platos típicos de una región) y realiza la petición HTTP si es necesario. "
        "Con el contexto recuperado, el historial de la conversación (memoria) y la consulta del usuario, se construye un prompt estructurado. "
        "Este prompt es enviado al LLM (Gemini), el cual genera una respuesta precisa y fundamentada, que finalmente se muestra al usuario en la interfaz."
    )

    add_heading(doc, "C. Gráficos de arquitectura y funcionamiento", 1)
    doc.add_picture(str(ARCH_PATH), width=Inches(6.2))
    doc.add_paragraph("Figura 1: Arquitectura de componentes del sistema PeruRecetas.")

    add_heading(doc, "D. Descripción de los pasos de configuración en el entorno de Python", 1)
    add_bullets(
        doc,
        [
            "1. Creación de un entorno virtual: Se recomienda usar 'python -m venv venv' para aislar las dependencias.",
            "2. Activación del entorno virtual: 'source venv/bin/activate' en Linux/Mac o 'venv\\Scripts\\activate' en Windows.",
            "3. Instalación de dependencias: Ejecutar 'pip install -r requirements.txt' para instalar Streamlit, FAISS, AIsuite, Google Generative AI, entre otros.",
            "4. Configuración de variables de entorno: Crear un archivo '.env' basado en '.env.example' y añadir la clave 'GOOGLE_API_KEY'.",
            "5. Ejecución de la aplicación: Iniciar la app web mediante el comando 'streamlit run app.py'.",
        ],
    )

    add_heading(doc, "E. Implementación del código en Python", 1)
    doc.add_paragraph(
        "La implementación principal reside en 'perurecetas_core.py'. Las características clave son:\n"
        "• Llamada al LLM: Se realiza a través de la clase AgenteRecetas en el método '_llamar_llm', que intenta usar AIsuite primero, y como respaldo (fallback) utiliza el SDK nativo 'google.generativeai'. "
        "El método 'enviar_mensaje' orquesta la construcción del prompt final.\n"
        "• Tools con Docstrings: Se implementaron funciones como 'consultar_api_comida_peru' (busca datos en la API) y 'recomendar_recetas_por_ingrediente' (sugiere opciones locales). Ambas cuentan con Docstrings detallados que explican sus argumentos y valores de retorno.\n"
        "• RAG: La clase 'PeruRecetasFAISS' maneja la indexación y recuperación. Se generan embeddings de los TXTs con Gemini y se indexan con FAISS ('IndexFlatIP'). En la recuperación, se devuelve el texto que hace 'match' semántico o textual (fallback) con la consulta del usuario."
    )

    add_heading(doc, "F. Descripción de la data utilizada para el RAG", 1)
    doc.add_paragraph("Se utilizaron al menos 5 recetarios (convertidos a TXT) especializados en cocina peruana y andina.")
    catalogo = catalogar_dataset()
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Archivo"
    hdr[1].text = "Ruta"
    hdr[2].text = "Tamano bytes"
    hdr[3].text = "Recetas"
    for item in catalogo:
        row = table.add_row().cells
        row[0].text = str(item["archivo"])
        row[1].text = str(item["ruta"])
        row[2].text = str(item["tamano_bytes"])
        row[3].text = str(item.get("recetas", "No calculado"))

    doc.save(DOCX_PATH)


def p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text.replace("\n", "<br/>"), style)


def crear_pdf() -> None:
    crear_diagrama()
    font_regular = "Helvetica"
    font_bold = "Helvetica-Bold"
    arial = Path("C:/Windows/Fonts/arial.ttf")
    arial_bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if arial.exists() and arial_bold.exists():
        pdfmetrics.registerFont(TTFont("Arial", str(arial)))
        pdfmetrics.registerFont(TTFont("Arial-Bold", str(arial_bold)))
        font_regular = "Arial"
        font_bold = "Arial-Bold"

    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font_regular
    styles.add(
        ParagraphStyle(
            name="TitleBlue",
            parent=styles["Title"],
            textColor=colors.HexColor("#0B2545"),
            fontName=font_bold,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H1Blue",
            parent=styles["Heading1"],
            textColor=colors.HexColor("#2E74B5"),
            fontName=font_bold,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            parent=styles["BodyText"],
            leading=14,
            spaceAfter=8,
            fontName=font_regular,
        )
    )

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )
    story = []
    story.append(p(PROJECT_TITLE, styles["TitleBlue"]))
    story.append(p(PROJECT_NAME, styles["Heading2"]))
    story.append(Spacer(1, 0.25 * inch))
    story.append(p(f"<b>Curso:</b> {CURSO}", styles["Body"]))
    story.append(p(f"<b>Docente:</b> {DOCENTE}", styles["Body"]))
    story.append(p("<b>Integrantes:</b><br/>" + "<br/>".join(INTEGRANTES), styles["Body"]))
    story.append(p(f"<b>Fecha de entrega:</b> {FECHA}", styles["Body"]))
    story.append(PageBreak())

    sections = [
        (
            "A. Descripción general de los modelos o librerías",
            "Gemini (LLM principal y embeddings), AIsuite (cliente unificado), FAISS (búsqueda semántica), Streamlit (interfaz web), y librerías estándar como requests para consultar la API."
        ),
        (
            "B. Flujo de trabajo personalizado",
            "El Agente recibe la consulta, busca contexto en FAISS (RAG), consulta la API externa si aplica, y combina la memoria con estos datos para formar un prompt. El LLM (Gemini) genera la respuesta basándose estrictamente en esta información estructurada."
        ),
    ]
    for title, body in sections:
        story.append(p(title, styles["H1Blue"]))
        story.append(p(body, styles["Body"]))

    story.append(p("C. Gráficos de arquitectura", styles["H1Blue"]))
    story.append(RLImage(str(ARCH_PATH), width=6.2 * inch, height=3.36 * inch))
    story.append(Spacer(1, 0.12 * inch))

    story.append(p("D. Configuración en el entorno de Python", styles["H1Blue"]))
    bullets = [
        "1. Crear y activar un entorno virtual.",
        "2. Instalar dependencias con 'pip install -r requirements.txt'.",
        "3. Configurar la GOOGLE_API_KEY en el archivo '.env'.",
        "4. Ejecutar la aplicación web con 'streamlit run app.py'."
    ]
    story.append(ListFlowable([ListItem(p(item, styles["Body"])) for item in bullets], bulletType="bullet"))

    story.append(p("E. Implementación del código en Python", styles["H1Blue"]))
    impl = [
        "Llamada al LLM: Implementada en 'AgenteRecetas._llamar_llm' utilizando AIsuite o google-generativeai.",
        "Tools con Docstrings: Funciones como 'consultar_api_comida_peru' documentadas con pydoc para explicar su propósito y parámetros.",
        "RAG: La clase 'PeruRecetasFAISS' realiza la vectorización con Gemini y recuperación con FAISS."
    ]
    story.append(ListFlowable([ListItem(p(item, styles["Body"])) for item in impl], bulletType="bullet"))

    story.append(p("F. Descripción de la data utilizada para el RAG", styles["H1Blue"]))
    story.append(p("Se utilizaron al menos 5 recetarios procesados en formato TXT estructurado.", styles["Body"]))
    catalogo = catalogar_dataset()
    rows = [["Archivo", "Ruta", "Tamano", "Recetas"]]
    for item in catalogo:
        rows.append([
            item["archivo"],
            item["ruta"],
            str(item["tamano_bytes"]),
            str(item.get("recetas", "No calculado")),
        ])
    table = Table(rows, colWidths=[1.4 * inch, 2.4 * inch, 1.2 * inch, 0.8 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C2CC")),
                ("FONTNAME", (0, 0), (-1, 0), font_bold),
                ("FONTNAME", (0, 1), (-1, -1), font_regular),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(table)

    doc.build(story)


def main() -> None:
    ensure_dirs()
    crear_docx()
    crear_pdf()
    print(f"DOCX generado: {DOCX_PATH}")
    print(f"PDF generado: {PDF_PATH}")


if __name__ == "__main__":
    main()
